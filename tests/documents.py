from io import BytesIO

from docx import Document
from pypdf import PdfWriter


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(lines: list[str]) -> bytes:
    body = ["BT", "/F1 12 Tf", "14 TL", "50 750 Td"]
    for line in lines:
        body += [f"({line}) Tj", "T*"]
    body.append("ET")
    stream = "\n".join(body)
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = b"%PDF-1.4\n"
    offsets = []
    for number, item in enumerate(objects, 1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n{item}\nendobj\n".encode("latin-1")
    xref = len(output)
    output += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    output += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    output += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return output


def make_blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
